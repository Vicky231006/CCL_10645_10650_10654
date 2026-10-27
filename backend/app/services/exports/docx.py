"""DOCX export using python-docx."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.schemas.resume import ResumeContent
from app.services.exports.links import ensure_url, link_label, mailto

LINK_COLOR = "1A4D8F"


def _add_hyperlink(paragraph, url: str, text: str, color: str = LINK_COLOR, underline: bool = True) -> None:
    """Append a real, clickable hyperlink run to a paragraph (python-docx has no API)."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


def render_docx(content: ResumeContent, page_size: str = "A4") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    section = doc.sections[0]
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(45)
    section.right_margin = Pt(45)

    contact = content.contact
    if contact.name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(contact.name)
        run.bold = True
        run.font.size = Pt(20)
    if contact.headline:
        p = doc.add_paragraph(contact.headline)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header_bits: list[tuple[str, str | None]] = []
    if contact.email:
        header_bits.append((contact.email, mailto(contact.email)))
    if contact.phone:
        header_bits.append((contact.phone, None))
    if contact.location:
        header_bits.append((contact.location, None))
    for link in contact.links:
        if link.url:
            header_bits.append((link.label or link_label(link.url), ensure_url(link.url) or None))
    if header_bits:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, (text, url) in enumerate(header_bits):
            if i:
                p.add_run(" • ")
            if url:
                _add_hyperlink(p, url, text)
            else:
                p.add_run(text)

    def heading(title: str):
        p = doc.add_paragraph()
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11.5)

    def bullet(text: str):
        doc.add_paragraph(text, style="List Bullet")

    order = content.section_order or ["summary", "experience", "projects", "education", "skills",
                                      "certifications", "achievements"]
    visibility = content.section_visibility or {}
    for sec in order:
        if visibility.get(sec) is False:
            continue
        if sec == "summary" and content.summary:
            heading("Summary")
            doc.add_paragraph(content.summary)
        elif sec == "experience" and content.experience:
            heading("Experience")
            for exp in content.experience:
                p = doc.add_paragraph()
                run = p.add_run((exp.role or "") + (f" — {exp.company}" if exp.company else ""))
                run.bold = True
                if exp.start or exp.end:
                    p.add_run(f"\t{exp.start} – {exp.end}")
                if exp.location:
                    doc.add_paragraph(exp.location)
                for b in exp.bullets:
                    bullet(b)
        elif sec == "projects" and content.projects:
            heading("Projects")
            for proj in content.projects:
                p = doc.add_paragraph()
                run = p.add_run(proj.name)
                run.bold = True
                if proj.tech:
                    p.add_run(f" | {', '.join(proj.tech)}")
                for b in proj.bullets:
                    bullet(b)
                if proj.link:
                    lp = doc.add_paragraph()
                    purl = ensure_url(proj.link)
                    if purl:
                        _add_hyperlink(lp, purl, link_label(proj.link, proj.link))
                    else:
                        lp.add_run(proj.link)
        elif sec == "education" and content.education:
            heading("Education")
            for ed in content.education:
                p = doc.add_paragraph()
                run = p.add_run(ed.institution + (f" — {ed.degree}" if ed.degree else ""))
                run.bold = True
                if ed.start or ed.end:
                    p.add_run(f"\t{ed.start} – {ed.end}")
                for d in ed.details:
                    bullet(d)
        elif sec == "skills" and content.skills:
            heading("Skills")
            for group in content.skills:
                if group.items:
                    label = f"{group.category}: " if group.category and group.category != "Skills" else ""
                    doc.add_paragraph(label + ", ".join(group.items))
        elif sec == "certifications" and content.certifications:
            heading("Certifications")
            for c in content.certifications:
                bullet(c)
        elif sec == "achievements" and content.achievements:
            heading("Achievements")
            for a in content.achievements:
                bullet(a)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
