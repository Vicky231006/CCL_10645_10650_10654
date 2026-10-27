"""Unit tests for the multi-format parser registry and structuring."""
from __future__ import annotations

from io import BytesIO

from app.services.parsers.base import detect_format
from app.services.parsers.registry import parse_document, registered_formats
from app.services.parsers.structuring import structure_text

SAMPLE = """Asha Rao
asha.rao@example.com | +91 90000 00000 | Bengaluru
github.com/asha

Summary
Backend-focused computer science student.

Education
IIT Example — B.Tech CSE
2022 - 2026

Experience
Software Intern — Acme
May 2025 - Jul 2025
• Built a REST API reducing latency by 30%.

Projects
Resumo — FastAPI, React
• Implemented a parser registry.

Skills
Languages: Python, TypeScript
"""


def test_registry_lists_expected_formats():
    formats = set(registered_formats())
    for expected in {"pdf", "docx", "txt", "md", "rtf", "odt", "tex", "html", "image", "doc"}:
        assert expected in formats


def test_detect_format_by_extension():
    assert detect_format("a.PDF", None, b"") == "pdf"
    assert detect_format("a.docx", None, b"") == "docx"
    assert detect_format("a.md", None, b"") == "md"


def test_detect_format_by_magic_bytes():
    assert detect_format("nofile", None, b"%PDF-1.7 ...") == "pdf"
    assert detect_format("nofile", None, b"{\\rtf1 hello") == "rtf"


def test_parse_txt_roundtrip():
    parsed = parse_document("txt", SAMPLE.encode())
    assert parsed.confidence > 0.8
    assert "Asha Rao" in parsed.text


def test_parse_markdown_strips_syntax():
    parsed = parse_document("md", b"# Heading\n**bold** and *italic*\n- item")
    assert "#" not in parsed.text
    assert "**" not in parsed.text


def test_parse_tex_does_not_execute_commands():
    tex = r"\documentclass{article}\begin{document}\section{Experience}\item Built an API\end{document}"
    parsed = parse_document("tex", tex.encode())
    assert "\\documentclass" not in parsed.text
    assert "Experience" in parsed.text


def test_parse_html_strips_scripts():
    html = b"<html><head><style>.a{}</style></head><body><script>alert(1)</script><p>Hello Resume</p></body></html>"
    parsed = parse_document("html", html)
    assert "Hello Resume" in parsed.text
    assert "alert" not in parsed.text


def test_parse_rtf():
    rtf = rb"{\rtf1\ansi Hello \b World\b0}"
    parsed = parse_document("rtf", rtf)
    assert "Hello" in parsed.text


def test_parse_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Asha Rao")
    doc.add_paragraph("asha@example.com")
    buf = BytesIO()
    doc.save(buf)
    parsed = parse_document("docx", buf.getvalue())
    assert "Asha Rao" in parsed.text
    assert parsed.confidence >= 0.8


def test_parse_pdf_image_only_reports_limitation():
    # Minimal invalid PDF-like content should not raise; returns low confidence.
    parsed = parse_document("pdf", b"%PDF-1.4 broken")
    assert parsed.confidence <= 0.2


def test_unsupported_format_is_honest():
    parsed = parse_document("xyz", b"data")
    assert parsed.confidence == 0.0
    assert parsed.warnings


def test_structuring_extracts_core_fields():
    content, confidence = structure_text(SAMPLE)
    assert content.contact.name == "Asha Rao"
    assert content.contact.email == "asha.rao@example.com"
    assert content.experience and content.experience[0].company.lower().endswith("acme")
    assert content.education and "iit" in content.education[0].institution.lower()
    assert content.skills and any("Python" in g.items for g in content.skills)
    assert confidence > 0.5


COMPACT = """Jordan Lee
jordan@example.com | +91 90000 00000 | Pune
Education
Example College of EngineeringSep. 2021 - June 2025
B.Tech in Computer Engineering CGPA: 9.1
City CollegeAug. 2019 - May 2021
Higher Secondary Percentage: 92
Experience
Research Fellow - IncomingOct. 2026 - July 2027
Example Foundation for IoT, IIT Pune Pune, India
- Selected for a funded fellowship working on
image-based detection for crops.
- Will conduct research.
Projects
Alpha App | Python, Flask Sep 2025
- Built a full-featured web application supporting
conversion and merging.
- Designed a responsive frontend.
Beta Tool|TypeScript|GitHub
- Engineered a playlist tool.
Hackathons & Competitions
CodeSprint | Winner 2025
"""


def test_structuring_splits_multiple_education_entries():
    content, _ = structure_text(COMPACT)
    assert len(content.education) == 2
    first = content.education[0]
    assert first.institution == "Example College of Engineering"
    assert first.start == "Sep. 2021" and first.end == "June 2025"
    assert "B.Tech" in first.degree


def test_structuring_keeps_company_line_with_role():
    content, _ = structure_text(COMPACT)
    assert len(content.experience) == 1
    exp = content.experience[0]
    assert "IIT Pune" in exp.company
    assert "Fellow" in exp.role
    assert exp.start == "Oct. 2026" and exp.end == "July 2027"
    assert len(exp.bullets) == 2  # wrapped continuation merged into its bullet


def test_structuring_merges_wrapped_bullets_and_tech():
    content, _ = structure_text(COMPACT)
    assert [p.name for p in content.projects] == ["Alpha App", "Beta Tool"]
    alpha = content.projects[0]
    assert alpha.tech == ["Python", "Flask"]
    assert len(alpha.bullets) == 2
    assert "conversion and merging" in alpha.bullets[0]
    beta = content.projects[1]
    assert "TypeScript" in beta.tech and beta.link == "GitHub"


def test_structuring_separates_achievements_heading():
    content, _ = structure_text(COMPACT)
    assert content.achievements == ["CodeSprint | Winner 2025"]
    assert content.summary == ""
