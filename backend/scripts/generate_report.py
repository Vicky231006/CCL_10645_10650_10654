"""Render REPORT.md (repo root) into REPORT.docx and REPORT.pdf.

The markdown uses a small, constrained subset so the academic report format can be
produced reproducibly:

  # / ## / ###        chapter / section / sub-section headings
  - item              bullet
  | a | b |           table (first row is the header; a --- separator row is skipped)
  %PAGE%              explicit page break
  %CENTER% ... %ENDCENTER%   centred block (each line one centred paragraph)
  %CODE% ... %ENDCODE%       monospaced preformatted block

Table captions ("**Table N ...**" placed just below the table in the source) are moved
above the table, matching the required format (table names above, figure names below).

Run from the backend/ directory with the project virtualenv:
    .\\.venv\\Scripts\\python.exe scripts\\generate_report.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
MD_PATH = ROOT / "REPORT.md"
DOCX_PATH = ROOT / "REPORT.docx"
PDF_PATH = ROOT / "REPORT.pdf"

SEP_RE = re.compile(r"^:?-{2,}:?$")


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def parse_table(rows: list[str]) -> list[list[str]]:
    parsed: list[list[str]] = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if cells and all(SEP_RE.match(c) for c in cells):
            continue  # separator row
        parsed.append(cells)
    return parsed


def parse_markdown(text: str):
    lines = text.splitlines()
    blocks: list[tuple[str, object]] = []
    mode = None
    buf: list[str] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if s == "%PAGE%":
            blocks.append(("page", None))
        elif s == "%CENTER%":
            mode, buf = "center", []
        elif s == "%ENDCENTER%":
            blocks.append(("center", buf))
            mode = None
        elif s == "%CODE%":
            mode, buf = "code", []
        elif s == "%ENDCODE%":
            blocks.append(("code", buf))
            mode = None
        elif mode == "center":
            buf.append(s)
        elif mode == "code":
            buf.append(raw)
        elif s.startswith("### "):
            blocks.append(("h3", s[4:].strip()))
        elif s.startswith("## "):
            blocks.append(("h2", s[3:].strip()))
        elif s.startswith("# "):
            blocks.append(("h1", s[2:].strip()))
        elif s.startswith("- "):
            blocks.append(("bullet", s[2:].strip()))
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(("table", parse_table(rows)))
            continue
        elif s.startswith("![") and "](" in s and s.endswith(")"):
            alt = s[2:s.find("](")]
            src = s[s.find("](") + 2:-1]
            blocks.append(("image", (alt, src)))
        elif s:
            blocks.append(("p", s))
        i += 1

    # Move a "**Table N ...**" caption that follows a table to above it.
    reordered: list[tuple[str, object]] = []
    for j, blk in enumerate(blocks):
        if (
            blk[0] == "table"
            and j + 1 < len(blocks)
            and blocks[j + 1][0] == "p"
            and str(blocks[j + 1][1]).lstrip().startswith("**Table")
        ):
            reordered.append(("caption", blocks[j + 1][1]))
            reordered.append(blk)
            blocks[j + 1] = ("_skip", None)
        elif blk[0] == "_skip":
            continue
        elif blk[0] == "p" and str(blk[1]).lstrip().startswith("**Figure"):
            reordered.append(("caption", blk[1]))
        else:
            reordered.append(blk)
    return reordered


INLINE_RE = re.compile(r"\*\*(.+?)\*\*")


def inline_parts(text: str) -> list[tuple[str, bool]]:
    """Split text into (segment, is_bold) parts, handling **bold** only."""
    parts: list[tuple[str, bool]] = []
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts or [(text, False)]


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def build_docx(blocks):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(2.54))

    # Footer page number.
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    frun = footer_p.add_run()
    frun._r.append(fld1); frun._r.append(instr); frun._r.append(fld2)
    frun.font.name = "Times New Roman"; frun.font.size = Pt(10)

    def para(text="", *, size=12, bold=False, italic=False, align=None,
             font="Times New Roman", space_after=6, space_before=0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(space_before)
        if align is not None:
            p.alignment = align
        if text:
            for seg, is_bold in inline_parts(text):
                r = p.add_run(seg)
                r.font.name = font
                r.font.size = Pt(size)
                r.bold = bold or is_bold
                r.italic = italic
        return p

    for kind, payload in blocks:
        if kind == "page":
            doc.add_page_break()
        elif kind == "h1":
            para(payload, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=6, space_after=14)
        elif kind == "h2":
            para(payload, size=12, bold=True, space_before=10, space_after=4)
        elif kind == "h3":
            para(payload, size=12, bold=True, italic=True, space_before=8, space_after=4)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            for seg, is_bold in inline_parts(payload):
                r = p.add_run(seg)
                r.font.name = "Times New Roman"; r.font.size = Pt(12); r.bold = is_bold
            p.paragraph_format.space_after = Pt(3)
        elif kind == "p":
            para(payload, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "caption":
            para(payload, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        elif kind == "center":
            for line in payload:
                if line.strip():
                    para(line.strip(), size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                         space_after=4)
                else:
                    para("", space_after=4)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(2)
            for idx, line in enumerate(payload):
                r = p.add_run(line)
                r.font.name = "Consolas"; r.font.size = Pt(8)
                if idx != len(payload) - 1:
                    r.add_break()
        elif kind == "image":
            alt, src = payload
            img_path = (ROOT / src) if not Path(src).is_absolute() else Path(src)
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                run.add_picture(str(img_path), width=Cm(15.8))
        elif kind == "table":
            table = payload
            if not table:
                continue
            ncols = max(len(r) for r in table)
            t = doc.add_table(rows=0, cols=ncols)
            t.style = "Table Grid"
            for ridx, row in enumerate(table):
                cells = t.add_row().cells
                for cidx in range(ncols):
                    txt = row[cidx] if cidx < len(row) else ""
                    cell_p = cells[cidx].paragraphs[0]
                    for seg, is_bold in inline_parts(txt):
                        r = cell_p.add_run(seg)
                        r.font.name = "Times New Roman"; r.font.size = Pt(8)
                        r.bold = (ridx == 0) or is_bold
            para("", space_after=2)

    doc.save(str(DOCX_PATH))


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def build_pdf(blocks):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Flowable, PageBreak, Paragraph, Preformatted,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    page_w, page_h = A4
    margin = 2.2 * cm
    avail = page_w - 2 * margin

    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=12, leading=15,
                          alignment=TA_JUSTIFY, spaceAfter=6)
    h1 = ParagraphStyle("h1", fontName="Times-Bold", fontSize=14, leading=18,
                        alignment=TA_CENTER, spaceBefore=6, spaceAfter=14)
    h2 = ParagraphStyle("h2", fontName="Times-Bold", fontSize=12, leading=15,
                        spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", fontName="Times-BoldItalic", fontSize=12, leading=15,
                        spaceBefore=8, spaceAfter=4)
    center = ParagraphStyle("center", fontName="Times-Roman", fontSize=12, leading=16,
                            alignment=TA_CENTER, spaceAfter=4)
    cap = ParagraphStyle("cap", fontName="Times-Roman", fontSize=10, leading=12,
                         alignment=TA_CENTER, spaceAfter=4)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=16, bulletIndent=4,
                            spaceAfter=3)
    code = ParagraphStyle("code", fontName="Courier", fontSize=8, leading=10,
                          spaceBefore=2, spaceAfter=6)

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def markup(t: str) -> str:
        out = []
        for seg, bold in inline_parts(t):
            out.append(f"<b>{esc(seg)}</b>" if bold else esc(seg))
        return "".join(out)

    class PageMarker(Flowable):
        def __init__(self, doc):
            super().__init__()
            self._doc = doc

        def wrap(self, *a):
            return (0, 0)

        def draw(self):
            self._doc.body = True
            self._doc.body_offset = self._doc.page - 1

    class ReportDoc(SimpleDocTemplate):
        def __init__(self, *args, **kw):
            super().__init__(*args, **kw)
            self.body = False
            self.body_offset = 0

        def afterPage(self):
            if self.page == 1:      # cover page is unnumbered
                return
            c = self.canv
            c.saveState()
            c.setFont("Times-Roman", 10)
            if self.body:
                label = str(self.page - self.body_offset)
            else:
                label = to_roman(self.page - 1)
            c.drawCentredString(page_w / 2.0, 1.3 * cm, label)
            c.restoreState()

    doc = ReportDoc(str(PDF_PATH), pagesize=A4, leftMargin=margin, rightMargin=margin,
                    topMargin=margin, bottomMargin=margin,
                    title="Resumo Mini-Project Report", author="Resumo Project Group")
    style_for = {2: 10, 3: 10, 4: 8, 5: 8, 6: 8, 7: 8, 8: 8}

    story: list = []
    for kind, payload in blocks:
        if kind == "page":
            story.append(PageBreak())
        elif kind == "h1":
            if str(payload).startswith("Chapter 1"):
                story.append(PageMarker(doc))
            story.append(Paragraph(markup(payload), h1))
        elif kind == "h2":
            story.append(Paragraph(markup(payload), h2))
        elif kind == "h3":
            story.append(Paragraph(markup(payload), h3))
        elif kind == "bullet":
            story.append(Paragraph(markup(payload), bullet, bulletText="\u2022"))
        elif kind == "p":
            story.append(Paragraph(markup(payload), body))
        elif kind == "caption":
            story.append(Paragraph(markup(payload), cap))
        elif kind == "center":
            for line in payload:
                story.append(Paragraph(markup(line.strip()) if line.strip() else "&nbsp;",
                                       center))
        elif kind == "code":
            story.append(Preformatted("\n".join(payload), code))
        elif kind == "image":
            alt, src = payload
            img_path = (ROOT / src) if not Path(src).is_absolute() else Path(src)
            if img_path.exists():
                from PIL import Image as PILImage
                from reportlab.platypus import Image as RLImage
                with PILImage.open(img_path) as im:
                    iw, ih = im.size
                aspect = ih / float(iw)
                target_w = avail
                target_h = target_w * aspect
                max_h = 17.0 * cm
                if target_h > max_h:
                    target_h = max_h
                    target_w = target_h / aspect
                story.append(Spacer(1, 6))
                story.append(RLImage(str(img_path), width=target_w, height=target_h))
                story.append(Spacer(1, 4))
        elif kind == "table":
            table = payload
            if not table:
                continue
            ncols = max(len(r) for r in table)
            fs = style_for.get(ncols, 8)
            cell = ParagraphStyle("cell", fontName="Times-Roman", fontSize=fs, leading=fs + 2)
            head = ParagraphStyle("head", parent=cell, fontName="Times-Bold")
            data = []
            for ridx, row in enumerate(table):
                r = []
                for cidx in range(ncols):
                    txt = row[cidx] if cidx < len(row) else ""
                    r.append(Paragraph(markup(txt),
                                       head if ridx == 0 else cell))
                data.append(r)
            tbl = Table(data, colWidths=[avail / ncols] * ncols, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 4))

    doc.build(story)


def to_roman(n: int) -> str:
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
            (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
            (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out or "i"


def main() -> int:
    if not MD_PATH.exists():
        print(f"ERROR: {MD_PATH} not found")
        return 1
    blocks = parse_markdown(MD_PATH.read_text(encoding="utf-8"))
    build_docx(blocks)
    build_pdf(blocks)
    print(f"Wrote {DOCX_PATH.name} and {PDF_PATH.name} ({len(blocks)} blocks)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
