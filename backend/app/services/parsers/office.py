"""Office format parsers: DOCX and legacy DOC."""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from app.services.parsers.base import ParsedDocument
from app.services.parsers.text_formats import parse_txt


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        from docx import Document
        from io import BytesIO

        doc = Document(BytesIO(data))
        lines: list[str] = []
        for para in doc.paragraphs:
            value = para.text.rstrip()
            if value:
                lines.append(value)
        # tables often hold contact/skills info in student resumes
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(dict.fromkeys(cells)))
        text = "\n".join(lines)
        if not text.strip():
            return ParsedDocument(
                text="",
                confidence=0.1,
                parser="docx",
                warnings=["DOCX contained no extractable text."],
            )
        return ParsedDocument(text=text, confidence=0.88, parser="docx")
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", confidence=0.0, parser="docx", warnings=[f"DOCX parse failed: {exc}"])


def parse_doc(data: bytes) -> ParsedDocument:
    """Legacy .doc requires an external converter. Never execute macros."""
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter:
        return ParsedDocument(
            text="",
            confidence=0.0,
            parser="doc",
            warnings=[
                "Legacy .doc conversion requires LibreOffice (soffice) which was not found. "
                "Please re-save the file as .docx or PDF and upload again."
            ],
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        try:
            subprocess.run(
                [converter, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=60,
            )
            out = Path(tmp) / "input.docx"
            if out.exists():
                from app.services.parsers.office import parse_docx as _parse_docx

                return _parse_docx(out.read_bytes())
        except Exception as exc:  # noqa: BLE001
            return ParsedDocument(text="", confidence=0.0, parser="doc", warnings=[f".doc conversion failed: {exc}"])
    return ParsedDocument(text="", confidence=0.0, parser="doc", warnings=[".doc conversion produced no output."])
